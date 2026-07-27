"""
Gerador de Monografia — Sistema de Gestão de Assistência Técnica Informática
Autor: Josimar Ferreira
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import datetime

doc = Document()

# ── Estilos de página ─────────────────────────────────────────────────────────
section = doc.sections[0]
section.page_width  = Inches(8.27)   # A4
section.page_height = Inches(11.69)
section.left_margin   = Cm(3.0)
section.right_margin  = Cm(2.0)
section.top_margin    = Cm(3.0)
section.bottom_margin = Cm(2.0)

# ── Helpers ───────────────────────────────────────────────────────────────────
def heading(text, level=1, bold=True, size=14, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    p.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    p.paragraph_format.space_after  = Pt(6)
    return p

def body(text, justified=True, size=12, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justified else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.first_line_indent = Cm(1.25)
    return p

def bullet(text, size=11):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(2)
    return p

def page_break():
    doc.add_page_break()

def add_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    # header row
    hdr = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = h
        cell.paragraphs[0].runs[0].bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(10)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), '1F4E79')
        tcPr.append(shd)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    # data rows
    for ri, row_data in enumerate(rows):
        row = table.rows[ri + 1]
        for ci, val in enumerate(row_data):
            cell = row.cells[ci]
            cell.text = str(val)
            cell.paragraphs[0].runs[0].font.size = Pt(10)
            if col_widths:
                cell.width = Inches(col_widths[ci])
    doc.add_paragraph()
    return table

# ═══════════════════════════════════════════════════════════════════════════════
#  CAPA
# ═══════════════════════════════════════════════════════════════════════════════
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(72)
r = p.add_run("REPÚBLICA DE ANGOLA")
r.bold = True; r.font.size = Pt(14)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("MINISTÉRIO DA EDUCAÇÃO")
r.bold = True; r.font.size = Pt(12)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("INSTITUTO SUPERIOR POLITÉCNICO")
r.bold = True; r.font.size = Pt(12)

doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("DESENVOLVIMENTO DE UM SISTEMA DE GESTÃO DE\nASSISTÊNCIA TÉCNICA INFORMÁTICA")
r.bold = True; r.font.size = Pt(18)
r.font.color.rgb = RGBColor(31, 78, 121)

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Monografia apresentada como requisito parcial para a obtenção\ndo grau de Licenciatura em Engenharia Informática")
r.font.size = Pt(12)

doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Autor:  Josimar Ferreira")
r.bold = True; r.font.size = Pt(13)

doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()
doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run(f"Luanda — {datetime.date.today().year}")
r.font.size = Pt(12)

page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  RESUMO
# ═══════════════════════════════════════════════════════════════════════════════
heading("RESUMO", level=1, size=14)
body(
    "O presente trabalho descreve o desenvolvimento de um Sistema de Gestão de Assistência "
    "Técnica Informática, concebido para automatizar e digitalizar os processos operacionais "
    "de uma empresa prestadora de serviços de manutenção e reparação de equipamentos "
    "informáticos. O sistema foi implementado com recurso às tecnologias PHP 8.2, MySQL 8.0, "
    "HTML5, CSS3, Bootstrap e JavaScript, seguindo o padrão arquitetural MVC (Model-View-Controller). "
    "A aplicação oferece três perfis de utilizador — Gerente, Técnico e Recepcionista — cada "
    "um com funcionalidades específicas e acesso controlado por sessões autenticadas. "
    "Os módulos implementados abrangem a gestão de clientes, equipamentos, orçamentos, "
    "serviços técnicos, stock de peças e componentes, compras, vendas, movimentação "
    "financeira, comissões de técnicos, relatórios e geração de documentos em PDF. "
    "Os resultados demonstram que o sistema satisfaz integralmente os requisitos levantados, "
    "reduzindo a dependência de processos manuais e aumentando a eficiência operacional da oficina."
)
doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("Palavras-chave: ")
r.bold = True; r.font.size = Pt(12)
r2 = p.add_run("Gestão de oficina, Assistência técnica informática, PHP, MySQL, MVC, Sistema de informação.")
r2.font.size = Pt(12)

page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  ABSTRACT
# ═══════════════════════════════════════════════════════════════════════════════
heading("ABSTRACT", level=1, size=14)
body(
    "This paper describes the development of an IT Technical Assistance Management System, "
    "designed to automate and digitise the operational processes of a company providing "
    "maintenance and repair services for computing equipment. The system was built using "
    "PHP 8.2, MySQL 8.0, HTML5, CSS3, Bootstrap and JavaScript, following the MVC "
    "(Model-View-Controller) architectural pattern. The application provides three user "
    "profiles — Manager, Technician and Receptionist — each with specific features and "
    "access controlled by authenticated sessions. The implemented modules cover client "
    "management, equipment tracking, budgeting, technical services, parts inventory, "
    "purchasing, sales, financial movement, technician commissions, reports and PDF document "
    "generation. The results show that the system fully meets the elicited requirements, "
    "reducing dependence on manual processes and increasing the operational efficiency of the workshop."
)
doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("Keywords: ")
r.bold = True; r.font.size = Pt(12)
r2 = p.add_run("Workshop management, IT technical assistance, PHP, MySQL, MVC, Information system.")
r2.font.size = Pt(12)

page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  ÍNDICE
# ═══════════════════════════════════════════════════════════════════════════════
heading("ÍNDICE GERAL", level=1, size=14)
toc_items = [
    ("1. FUNDAMENTAÇÃO TEÓRICA", ""),
    ("   1.1 Sistemas de Informação e Gestão Empresarial", ""),
    ("   1.2 Assistência Técnica Informática — Contexto e Desafios", ""),
    ("   1.3 Engenharia de Software — Conceitos Fundamentais", ""),
    ("   1.4 Padrão Arquitetural MVC", ""),
    ("   1.5 Tecnologias Utilizadas", ""),
    ("2. METODOLOGIA", ""),
    ("   2.1 Abordagem de Desenvolvimento", ""),
    ("   2.2 Levantamento de Requisitos", ""),
    ("   2.3 Ferramentas e Ambiente de Desenvolvimento", ""),
    ("3. RESULTADOS OBTIDOS", ""),
    ("   3.1 Requisitos Funcionais", ""),
    ("   3.2 Requisitos Não Funcionais", ""),
    ("   3.3 Arquitectura do Sistema", ""),
    ("   3.4 Diagrama de Casos de Uso", ""),
    ("   3.5 Diagrama de Classes", ""),
    ("   3.6 Diagrama Entidade-Relacionamento (ER)", ""),
    ("   3.7 Diagrama de Sequência", ""),
    ("   3.8 Diagrama de Actividades", ""),
    ("   3.9 Funcionalidades por Perfil de Utilizador", ""),
    ("       3.9.1 Perfil Gerente (adimin)", ""),
    ("       3.9.2 Perfil Técnico (tecnico)", ""),
    ("       3.9.3 Perfil Recepcionista (recep)", ""),
    ("4. CONCLUSÕES E TRABALHO FUTURO", ""),
    ("REFERÊNCIAS BIBLIOGRÁFICAS", ""),
]
for item, _ in toc_items:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(item)
    r.font.size = Pt(11)

page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  INTRODUÇÃO
# ═══════════════════════════════════════════════════════════════════════════════
heading("INTRODUÇÃO", level=1, size=14)
body(
    "A rápida proliferação de equipamentos informáticos nas últimas décadas — computadores, "
    "portáteis, impressoras, servidores, tablets e telemóveis — gerou uma procura crescente "
    "por serviços especializados de manutenção e reparação. Em Angola, este sector tem crescido "
    "de forma acelerada, mas a maioria das empresas de assistência técnica ainda opera com base "
    "em processos manuais: registos em papel, orçamentos manuscritos e controlo de stock "
    "informal. Esta realidade limita a capacidade de atendimento, dificulta o controlo "
    "financeiro e compromete a qualidade do serviço prestado ao cliente."
)
body(
    "O presente trabalho surge da necessidade de desenvolver uma solução informática que "
    "digitalize e automatize os processos de uma oficina de assistência técnica informática, "
    "nomeadamente o registo de clientes e equipamentos, a gestão de orçamentos e ordens de "
    "serviço, o controlo de stock de peças, o fluxo financeiro e a geração de relatórios. "
    "O sistema foi concebido para responder ao contexto angolano, com interface em português, "
    "suporte à moeda kwanza (Kz) e configurações adaptáveis por variáveis de ambiente."
)
body(
    "O trabalho está organizado em três capítulos principais: a Fundamentação Teórica, que "
    "estabelece os conceitos base; a Metodologia, que descreve o processo de desenvolvimento; "
    "e os Resultados Obtidos, que apresenta em detalhe a arquitectura, os diagramas UML e as "
    "funcionalidades implementadas para cada perfil de utilizador."
)
page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  CAP 1 — FUNDAMENTAÇÃO TEÓRICA
# ═══════════════════════════════════════════════════════════════════════════════
heading("1. FUNDAMENTAÇÃO TEÓRICA", level=1, size=14)

heading("1.1 Sistemas de Informação e Gestão Empresarial", level=2, size=13)
body(
    "Um sistema de informação (SI) é definido como um conjunto organizado de recursos — "
    "humanos, tecnológicos e organizacionais — que recolhe, processa, armazena e comunica "
    "informação para suportar a tomada de decisão, a coordenação e o controlo de uma "
    "organização (Laudon & Laudon, 2021). Os sistemas de informação de gestão (SIG) "
    "constituem uma categoria especial de SI orientada para fornecer aos gestores informação "
    "estruturada e resumida sobre as operações da empresa."
)
body(
    "Para as pequenas e médias empresas (PME), os SIG representam um instrumento fundamental "
    "de competitividade: permitem reduzir erros operacionais, acelerar o atendimento ao cliente "
    "e fornecer dados em tempo real sobre o desempenho do negócio. Turban et al. (2018) "
    "destacam que a digitalização dos processos operacionais através de sistemas de informação "
    "é um dos principais factores de diferenciação competitiva no século XXI."
)

heading("1.2 Assistência Técnica Informática — Contexto e Desafios", level=2, size=13)
body(
    "A assistência técnica informática engloba serviços de diagnóstico, manutenção preventiva, "
    "manutenção correctiva e reparação de equipamentos de tecnologia da informação. A gestão "
    "eficaz destes serviços exige o controlo coordenado de múltiplos processos: entrada de "
    "equipamentos, diagnóstico técnico, elaboração de orçamentos, aquisição de peças, "
    "execução da reparação e entrega ao cliente."
)
body(
    "Os principais desafios identificados no contexto das oficinas de informática em Angola incluem: "
    "(i) ausência de rastreabilidade dos equipamentos em reparação; "
    "(ii) dificuldade na gestão do stock de componentes; "
    "(iii) falta de transparência no cálculo de orçamentos e comissões; "
    "(iv) inexistência de dados históricos que permitam análises financeiras; "
    "(v) comunicação deficiente entre técnicos e recepção. "
    "Um sistema de gestão integrado é a solução natural para estes problemas."
)

heading("1.3 Engenharia de Software — Conceitos Fundamentais", level=2, size=13)
body(
    "A Engenharia de Software (ES) é a área da informática que estuda a aplicação sistemática "
    "e disciplinada de princípios de engenharia ao desenvolvimento, operação e manutenção de "
    "software (Sommerville, 2019). Os seus pilares incluem a especificação de requisitos, "
    "o design de arquitectura, a implementação, os testes e a manutenção."
)
body(
    "Para o presente projecto foi seguido um ciclo de vida incremental e iterativo, no qual "
    "o sistema foi desenvolvido em incrementos funcionais. Esta abordagem permitiu validar "
    "progressivamente as funcionalidades e incorporar ajustes ao longo do desenvolvimento, "
    "minimizando o risco de desvio face aos requisitos reais."
)
body(
    "A UML (Unified Modeling Language) foi utilizada como notação padrão para a modelação do "
    "sistema, conforme recomendado pelo OMG (Object Management Group). Os diagramas "
    "elaborados — casos de uso, classes, entidade-relacionamento, sequência e actividades — "
    "são apresentados no Capítulo 3."
)

heading("1.4 Padrão Arquitetural MVC", level=2, size=13)
body(
    "O padrão Model-View-Controller (MVC) é um padrão de arquitectura de software que separa "
    "uma aplicação em três componentes lógicos (Gamma et al., 1994):"
)
bullet("Model (Modelo): responsável pela lógica de negócio e pelo acesso aos dados. "
       "No sistema desenvolvido, os modelos encontram-se em app/adms/Models/ e utilizam PDO "
       "para comunicar com a base de dados MySQL.")
bullet("View (Vista): responsável pela apresentação da interface ao utilizador. "
       "As vistas estão em app/adms/Views/ e são ficheiros PHP que geram HTML dinâmico "
       "com Bootstrap 4.")
bullet("Controller (Controlador): responsável por receber as acções do utilizador, "
       "invocar o modelo adequado e seleccionar a vista a renderizar. "
       "Os controladores residem em app/adms/Controllers/.")
body(
    "A separação de responsabilidades proporcionada pelo MVC facilita a manutenção, "
    "os testes e a evolução do sistema, pois alterações na interface não afectam a lógica "
    "de negócio e vice-versa. O roteamento é realizado pelo ConfigController, que interpreta "
    "o parâmetro GET 'url' e instancia dinamicamente o controlador correspondente."
)

heading("1.5 Tecnologias Utilizadas", level=2, size=13)
body(
    "A selecção das tecnologias foi orientada pelos critérios de maturidade, disponibilidade "
    "de documentação, custo zero de licenciamento e compatibilidade com o ambiente XAMPP "
    "(servidor local amplamente utilizado no contexto angolano)."
)
add_table(
    ["Tecnologia / Ferramenta", "Versão", "Função no Sistema"],
    [
        ["PHP", "8.2", "Linguagem de backend — lógica de negócio e controlo MVC"],
        ["MySQL", "8.0+", "Sistema de gestão de base de dados relacional"],
        ["HTML5 / CSS3", "—", "Estrutura e estilo das páginas web"],
        ["Bootstrap", "4.x", "Framework CSS responsivo para interface gráfica"],
        ["JavaScript / jQuery", "3.x", "Interactividade do lado do cliente"],
        ["PHPMailer", "6.12", "Envio de e-mails (recuperação de senha, notificações)"],
        ["mPDF", "8.3", "Geração de documentos PDF (relatórios, orçamentos)"],
        ["Composer", "2.x", "Gestão de dependências PHP (PSR-4 autoloading)"],
        ["XAMPP", "8.2", "Ambiente de desenvolvimento local (Apache + MySQL + PHP)"],
        ["Git", "—", "Controlo de versões do código-fonte"],
    ]
)

page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  CAP 2 — METODOLOGIA
# ═══════════════════════════════════════════════════════════════════════════════
heading("2. METODOLOGIA", level=1, size=14)

heading("2.1 Abordagem de Desenvolvimento", level=2, size=13)
body(
    "O desenvolvimento do sistema seguiu uma abordagem incremental e iterativa, inspirada "
    "nos princípios das metodologias ágeis. Em vez de definir todos os requisitos à partida "
    "e desenvolver o sistema de uma só vez (modelo em cascata), optou-se por construir o "
    "sistema em incrementos funcionais testáveis, cada um acrescentando novos módulos "
    "e funcionalidades ao núcleo já validado."
)
body(
    "Esta escolha justifica-se pelo facto de o cliente (proprietário da oficina) não ter "
    "experiência prévia com sistemas informáticos, tornando difícil a especificação completa "
    "dos requisitos no início. A abordagem iterativa permitiu apresentar versões intermédias "
    "do sistema, recolher feedback e ajustar as funcionalidades antes de avançar para o "
    "módulo seguinte."
)
body(
    "O desenvolvimento decorreu em quatro iterações principais: (1) módulo de autenticação "
    "e gestão de utilizadores; (2) módulo de clientes, equipamentos e orçamentos; "
    "(3) módulo de stock, compras, vendas e finanças; (4) módulo de relatórios, gráficos "
    "e geração de PDF."
)

heading("2.2 Levantamento de Requisitos", level=2, size=13)
body(
    "O levantamento de requisitos foi realizado através de três técnicas complementares:"
)
bullet("Entrevistas semi-estruturadas com o gestor da oficina, para compreender os processos "
       "actuais, os problemas recorrentes e as expectativas face ao sistema.")
bullet("Observação directa do fluxo de trabalho da oficina, permitindo identificar "
       "etapas não verbalizadas nas entrevistas (e.g., o processo informal de registo de "
       "entrada de equipamentos).")
bullet("Análise de documentos existentes: recibos manuscritos, listas de preços e "
       "folhas de controlo de stock, que serviram de base para definir os campos de cada módulo.")
body(
    "Os requisitos foram classificados em funcionais (o que o sistema deve fazer) e não "
    "funcionais (como o sistema deve comportar-se), conforme apresentado no Capítulo 3."
)

heading("2.3 Ferramentas e Ambiente de Desenvolvimento", level=2, size=13)
body(
    "O ambiente de desenvolvimento foi configurado com XAMPP 8.2 sobre Windows 10, "
    "proporcionando um servidor Apache local, servidor MySQL e interpretador PHP. "
    "O editor de código utilizado foi o NetBeans IDE (evidenciado pela presença da pasta "
    "nbproject/ no repositório), complementado pelo Visual Studio Code para edição de "
    "ficheiros de configuração e scripts. O controlo de versões foi gerido com Git, "
    "com repositório local. A gestão de dependências PHP foi realizada com Composer, "
    "que instala automaticamente as bibliotecas PHPMailer e mPDF definidas no composer.json."
)
body(
    "Para a modelação do sistema foram utilizados diagramas UML, elaborados com recurso "
    "a ferramentas de modelação e apresentados nas secções seguintes através de descrição "
    "textual estruturada, dado o contexto de produção deste documento."
)

page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  CAP 3 — RESULTADOS OBTIDOS
# ═══════════════════════════════════════════════════════════════════════════════
heading("3. RESULTADOS OBTIDOS", level=1, size=14)
body(
    "Este capítulo apresenta os artefactos técnicos produzidos ao longo do desenvolvimento, "
    "incluindo os requisitos do sistema, a sua arquitectura, os principais diagramas UML "
    "e a descrição detalhada das funcionalidades disponíveis para cada perfil de utilizador."
)

# ── 3.1 Requisitos Funcionais ──────────────────────────────────────────────────
heading("3.1 Requisitos Funcionais", level=2, size=13)
body(
    "Os requisitos funcionais descrevem as funcionalidades que o sistema deve disponibilizar. "
    "Foram organizados por módulo, conforme a tabela seguinte."
)
add_table(
    ["ID", "Módulo", "Descrição do Requisito"],
    [
        ["RF01", "Autenticação", "O sistema deve permitir o login com e-mail e senha, diferenciando os perfis adimin, tecnico e recep."],
        ["RF02", "Autenticação", "O sistema deve suportar recuperação de senha por e-mail, com código de verificação de 6 dígitos com expiração."],
        ["RF03", "Utilizadores", "O gerente deve poder criar, editar e desactivar contas de técnicos e recepcionistas."],
        ["RF04", "Clientes", "O sistema deve permitir o registo, edição e consulta de clientes com NBI, NIF, contacto e morada."],
        ["RF05", "Equipamentos", "O sistema deve registar equipamentos informáticos (PC, portátil, impressora, servidor, tablet, telemóvel, etc.) associados a clientes."],
        ["RF06", "Equipamentos", "O sistema deve registar o estado do equipamento (Recebido, Em Diagnóstico, Aguardando Peças, Em Reparação, Concluído, Entregue)."],
        ["RF07", "Orçamentos", "O sistema deve permitir criar orçamentos / ordens de serviço, associando equipamento, técnico e tipo de serviço."],
        ["RF08", "Orçamentos", "O sistema deve permitir adicionar peças ao orçamento, actualizando o stock automaticamente."],
        ["RF09", "Orçamentos", "O sistema deve suportar desconto configurável sobre o valor do orçamento."],
        ["RF10", "Orçamentos", "O sistema deve gerar o orçamento em PDF para impressão ou envio ao cliente."],
        ["RF11", "Serviços", "O técnico deve poder consultar e actualizar o estado dos serviços que lhe estão atribuídos."],
        ["RF12", "Stock", "O sistema deve manter o stock de peças e componentes, alertando quando atingir o nível mínimo configurável."],
        ["RF13", "Compras", "O sistema deve registar compras a fornecedores, actualizando o stock e criando a conta a pagar correspondente."],
        ["RF14", "Vendas", "O sistema deve registar vendas de peças, actualizando o stock e registando na movimentação financeira."],
        ["RF15", "Finanças", "O sistema deve registar toda a movimentação de caixa (entradas e saídas) e calcular saldos do dia e do mês."],
        ["RF16", "Contas", "O sistema deve gerir contas a pagar e contas a receber, com controlo do estado de pagamento."],
        ["RF17", "Comissões", "O sistema deve calcular e registar automaticamente a comissão do técnico (percentagem configurável) por serviço concluído."],
        ["RF18", "Relatórios", "O sistema deve gerar relatórios de serviços, orçamentos, movimentação, compras, vendas, contas e comissões, filtráveis por período e estado."],
        ["RF19", "Gráficos", "O sistema deve apresentar gráficos de desempenho financeiro ao gerente."],
        ["RF20", "Chat", "O sistema deve dispor de uma funcionalidade de chat interno entre utilizadores."],
        ["RF21", "Perfil", "Cada utilizador deve poder editar o seu perfil e foto."],
        ["RF22", "Entrada", "O sistema deve registar a entrada de equipamentos na oficina com data, técnico responsável e serviço a realizar."],
    ],
    col_widths=[0.5, 1.3, 4.5]
)

# ── 3.2 Requisitos Não Funcionais ─────────────────────────────────────────────
heading("3.2 Requisitos Não Funcionais", level=2, size=13)
add_table(
    ["ID", "Categoria", "Descrição"],
    [
        ["RNF01", "Segurança", "As senhas dos utilizadores são armazenadas com hash MD5; em produção recomenda-se bcrypt/Argon2."],
        ["RNF02", "Segurança", "As páginas restritas verificam a sessão PHP antes de renderizar qualquer conteúdo."],
        ["RNF03", "Segurança", "O acesso directo a ficheiros PHP internos é bloqueado pela constante R4F5CC verificada em todos os ficheiros."],
        ["RNF04", "Desempenho", "O sistema deve responder a qualquer acção do utilizador em menos de 3 segundos, em condições de rede local."],
        ["RNF05", "Usabilidade", "A interface deve ser responsiva (Bootstrap), adaptando-se a ecrãs de desktop e tablet."],
        ["RNF06", "Manutenibilidade", "O código deve seguir o padrão PSR-4 de autoloading, com separação clara entre camadas MVC."],
        ["RNF07", "Portabilidade", "O sistema deve funcionar em qualquer servidor com PHP 8.x, MySQL 8.x e servidor web Apache ou Nginx."],
        ["RNF08", "Configurabilidade", "Todos os parâmetros de negócio (comissão, desconto, stock mínimo) devem ser configuráveis via ficheiro .env, sem alterar código."],
        ["RNF09", "Disponibilidade", "Para uso em produção, o sistema deve estar disponível 24/7, com backups automáticos diários da base de dados."],
        ["RNF10", "Internacionalização", "A interface e os documentos gerados devem estar em língua portuguesa, com valores monetários em Kwanza (Kz)."],
    ],
    col_widths=[0.6, 1.4, 4.3]
)

# ── 3.3 Arquitectura ──────────────────────────────────────────────────────────
heading("3.3 Arquitectura do Sistema", level=2, size=13)
body(
    "O sistema adopta a arquitectura em três camadas (Three-Tier Architecture) combinada com "
    "o padrão MVC, conforme ilustrado na descrição estrutural abaixo:"
)

arch_text = """
┌─────────────────────────────────────────────────────────────┐
│                    CAMADA DE APRESENTAÇÃO                    │
│  (Views — app/adms/Views/)                                   │
│  HTML5 · CSS3 · Bootstrap 4 · JavaScript / jQuery           │
├─────────────────────────────────────────────────────────────┤
│                    CAMADA DE NEGÓCIO                         │
│  (Controllers — app/adms/Controllers/)                       │
│  ConfigController (router) · Permissao (ACL) · Controllers   │
│                                                             │
│  (Models — app/adms/Models/)                                 │
│  AdmsHome · AdmsTecnico · AdmsRecepcionista                  │
│  AdmsLogin · AdmsPerfil · AdmsGraficos · Conn (PDO)          │
├─────────────────────────────────────────────────────────────┤
│                    CAMADA DE DADOS                           │
│  MySQL 8.0  ·  Base de dados: manutencao                    │
│  19 tabelas · 6 views · InnoDB · utf8mb4                    │
└─────────────────────────────────────────────────────────────┘

  UTILITÁRIOS TRANSVERSAIS
  PHPMailer 6.12 (e-mail)  ·  mPDF 8.3 (PDF)
  Config.php (.env loader)  ·  Composer (PSR-4 autoload)
"""
p = doc.add_paragraph()
run = p.add_run(arch_text)
run.font.name = "Courier New"
run.font.size = Pt(9)

body(
    "O fluxo de uma requisição HTTP típica é o seguinte: o utilizador acede a uma URL do tipo "
    "index.php?url=orcamento; o index.php inicia a sessão e instancia o ConfigController; "
    "o ConfigController carrega as constantes de configuração, verifica as permissões via "
    "classe Permissao e instancia dinamicamente o controlador Orcamento; o controlador "
    "Orcamento invoca o modelo AdmsTecnico para obter/persistir dados; os dados são passados "
    "à ConfigView, que inclui o ficheiro de vista correspondente e renderiza o HTML final."
)

# ── 3.4 Diagrama de Casos de Uso ──────────────────────────────────────────────
heading("3.4 Diagrama de Casos de Uso", level=2, size=13)
body(
    "O diagrama de casos de uso representa as interacções entre os actores (utilizadores) "
    "e o sistema. Os três actores principais são o Gerente (adimin), o Técnico (tecnico) "
    "e a Recepcionista (recep). O actor Sistema representa processos automáticos."
)

uc_text = """
                    ┌──────────────────────────────────────────┐
                    │        SISTEMA DE GESTÃO ATI             │
  ┌─────────┐       │                                          │
  │ GERENTE │──────►│ UC01 Gerir Utilizadores                  │
  │ (adimin)│       │ UC02 Gerir Clientes                      │
  └─────────┘       │ UC03 Gerir Equipamentos                  │
       │            │ UC04 Gerir Fornecedores                  │
       │            │ UC05 Gerir Produtos / Stock              │
       │            │ UC06 Gerir Tipo de Serviço               │
       │            │ UC07 Gerir Orçamentos / OS               │
       │            │ UC08 Ver Movimentação Financeira         │
       │            │ UC09 Gerir Compras                       │
       │            │ UC10 Gerir Vendas                        │
       │            │ UC11 Consultar Estoque Baixo             │
       │            │ UC12 Visualizar Gráficos                 │
       │            │ UC13 Gerar Relatórios                    │
  ┌──────────────┐  │ UC14 Consultar Comissões                 │
  │   TÉCNICO    │──►│ UC15 Consultar OS Atribuídas            │
  │  (tecnico)   │  │ UC16 Actualizar Estado Serviço           │
  └──────────────┘  │ UC17 Ver Comissões Próprias             │
       │            │ UC18 Gerar Relatório Comissão            │
  ┌───────────────┐ │ UC19 Gerir Contas a Pagar               │
  │RECEPCIONISTA  │─►│ UC20 Gerir Contas a Receber             │
  │   (recep)    │  │ UC21 Registar Entrada Equipamento       │
  └───────────────┘ │ UC22 Criar / Aprovar Orçamento          │
       │            │ UC23 Chat Interno                        │
  (todos)           │ UC24 Gerir Perfil / Foto                 │
       └───────────►│ UC25 Login / Logout                      │
                    │ UC26 Recuperar Senha                     │
                    └──────────────────────────────────────────┘
"""
p = doc.add_paragraph()
p.add_run(uc_text).font.name = "Courier New"
p.runs[0].font.size = Pt(8)

# ── 3.5 Diagrama de Classes ───────────────────────────────────────────────────
heading("3.5 Diagrama de Classes", level=2, size=13)
body(
    "O diagrama de classes reflecte a estrutura orientada a objectos do backend do sistema, "
    "organizado em três pacotes principais: Core, Controllers e Models."
)

cls_text = """
 ┌──────────────────────────────────┐
 │  <<package>> Core                │
 │                                  │
 │  ┌────────────┐                  │
 │  │   Config   │                  │
 │  │────────────│                  │
 │  │-config[]   │                  │
 │  │+load()     │                  │
 │  │+get()      │                  │
 │  └─────┬──────┘                  │
 │        │uses                     │
 │  ┌─────▼──────────────────────┐  │
 │  │   ConfigController         │  │
 │  │────────────────────────────│  │
 │  │-url: string                │  │
 │  │+carregar()                 │  │
 │  │-config()                   │  │
 │  └─────┬──────────────────────┘  │
 │        │uses                     │
 │  ┌─────▼──────────┐              │
 │  │   Permissao    │              │
 │  │────────────────│              │
 │  │-pgPublica[]    │              │
 │  │-pgRestrita[]   │              │
 │  │+index()        │              │
 │  └────────────────┘              │
 │  ┌─────────────────┐             │
 │  │  ConfigView     │             │
 │  │─────────────────│             │
 │  │-view: string    │             │
 │  │+renderizar()    │             │
 │  │+renderizarLogin()│            │
 │  │+renderizaRelatorio()│         │
 │  └─────────────────┘             │
 └──────────────────────────────────┘

 ┌────────────────────────────────────────┐
 │  <<package>> Models                    │
 │                                        │
 │  ┌─────────────┐                       │
 │  │    Conn     │  (classe base)         │
 │  │─────────────│                       │
 │  │-connect:PDO │                       │
 │  │+connect()   │                       │
 │  └──────┬──────┘                       │
 │         │ herda                        │
 │  ┌──────┴──────┬───────────────────┐   │
 │  │             │                   │   │
 │  ▼             ▼                   ▼   │
 │ AdmsHome  AdmsTecnico  AdmsRecepcionista│
 │ AdmsLogin AdmsPerfil  AdmsGraficos     │
 │ AdmsMpdf  AdmsRecuperarSenha           │
 └────────────────────────────────────────┘

 ┌──────────────────────────────────────────┐
 │  <<package>> Controllers                 │
 │                                          │
 │  Login · Home · Dashboard · Perfil       │
 │  Cliente · Equipamento · Fornecedor      │
 │  Orcamento · OrcamentoRecepcao           │
 │  AddProdutoOrcamento · Servico           │
 │  Produto · Categoria · Estoque           │
 │  TipoServico · Compras · Vendas          │
 │  ContasPagar · ContaReceber              │
 │  Movimentacao · Comissoes                │
 │  Relatorio · RelatorioTecnico            │
 │  Consultas · Graficos · Chat             │
 │  EntradaEquipamento · RecuperarSenha     │
 │  Tecnico · Recepcionista · EditarFoto    │
 └──────────────────────────────────────────┘
"""
p = doc.add_paragraph()
p.add_run(cls_text).font.name = "Courier New"
p.runs[0].font.size = Pt(8)

# ── 3.6 Diagrama ER ────────────────────────────────────────────────────────────
heading("3.6 Diagrama Entidade-Relacionamento (ER)", level=2, size=13)
body(
    "A base de dados do sistema contém 19 tabelas e 6 views, organizadas no diagrama "
    "entidade-relacionamento que se segue. As entidades principais e as suas relações "
    "são descritas textualmente e depois tabeladas."
)

er_text = """
 usuario (1) ─────────────────────────────────── (N) control_usuario
     │
     │  (utilizadores do sistema: adimin / tecnico / recep)

 clientes (1) ──────────────── (N) equipamento
     │                               │
     │                               │ (N)
     │                           orcamentos (N) ─────── (N) orc_prod
     │                               │                        │
     │                               │                    produto (N) ──── (1) categoria
     │                               │                        │
     │                           tipo_servico            fornecedor (1) ─── (N) compras
     │
     └── conntas_areceber (ligada a orcamentos)

 tecnicos (1) ──── (N) orcamentos  (campo tecnico = NIF)
 tecnicos (1) ──── (N) comissao    (campo niftecnico)
 tecnicos (1) ──── (N) entrada_equipamento

 contas_apagar (1) ──── (N) compras
 movimentacao  (independente — registo de caixa)
 vendas        (independente — registo de venda de peças)
 reset_senha   (ligada a usuario por email)
"""
p = doc.add_paragraph()
p.add_run(er_text).font.name = "Courier New"
p.runs[0].font.size = Pt(9)

body("A tabela seguinte resume as entidades, os seus atributos principais e as cardinalidades:")
add_table(
    ["Entidade", "Atributos Chave", "Relações"],
    [
        ["usuario", "idusuario, nome, email, senha, nivel (adimin/tecnico/recep)", "1:N com control_usuario"],
        ["clientes", "idclientes, nbi, nif, nome, sobrenome, email, telefone, morada", "1:N com equipamento; 1:N com conntas_areceber"],
        ["equipamento", "idequipamento, idcliente, numero_serie, tipo_equipamento, marca, modelo, estado", "N:1 com clientes; 1:N com orcamentos"],
        ["tecnicos", "idtecnico, nbi, nif, nome, sobrenome, email, telefone", "1:N com orcamentos, comissao, entrada_equipamento"],
        ["recepcionista", "idrecepcionista, nbi, nif, nome, sobrenome, email, telefone", "Registo independente"],
        ["orcamentos", "idorcamentos, veiculo (numero_serie), id_tipo_servico, valor, data, tecnico, status, tipo", "N:1 equipamento; N:1 tipo_servico; 1:N orc_prod; 1:N conntas_areceber"],
        ["orc_prod", "idorc_prod, orcamentos, produtos, quantidade", "N:1 orcamentos; N:1 produto"],
        ["produto", "idproduto, idfornecedor, idcategoria, nome, valor_compra, valor_venda, estoque", "N:1 fornecedor; N:1 categoria; 1:N orc_prod; 1:N compras"],
        ["fornecedor", "idfornecedor, nbi, nif, nome, email, telefone, tipo_pessoa", "1:N produto; 1:N compras"],
        ["categoria", "idcategoria, nome", "1:N produto"],
        ["tipo_servico", "idtipo_servico, nome, valor", "1:N orcamentos"],
        ["contas_apagar", "idcontas_apagar, descricao, valor, data_venci, pago", "1:N compras"],
        ["conntas_areceber", "idconntas_areceber, idorcamentos, valortotal, adiantameto, pago", "N:1 orcamentos"],
        ["movimentacao", "id, tipo (Entrada/Saída), descricao, valor, funcionario, data", "Independente"],
        ["compras", "idcompras, idproduto, idcontas_apagar, valor, quantidade_estoque", "N:1 produto; N:1 contas_apagar"],
        ["vendas", "idvendas, produto, valor, funcionario, data", "Independente"],
        ["comissao", "id, valor, servico, tipo, data, niftecnico", "N:1 tecnicos (por NIF)"],
        ["entrada_equipamento", "id, numero_serie, tipo_equipamento, cliente, niftecnico, estado, data_entrada", "N:1 tecnicos (por NIF)"],
        ["reset_senha", "id, email, codigo, expira_em, usado", "N:1 usuario (por email)"],
    ],
    col_widths=[1.3, 2.5, 2.5]
)

# ── 3.7 Diagrama de Sequência ─────────────────────────────────────────────────
heading("3.7 Diagrama de Sequência — Criação de Orçamento", level=2, size=13)
body(
    "O diagrama de sequência abaixo descreve o fluxo de mensagens entre os componentes "
    "do sistema durante a criação de um orçamento pelo Técnico ou Gerente."
)
seq_text = """
 Utilizador      Browser       index.php    ConfigController   Permissao   Orcamento(Ctrl)  AdmsTecnico(Model)   BD MySQL
     │              │              │               │               │              │                 │               │
     │──[POST]──────►│             │               │               │              │                 │               │
     │  btnAbriOrcamento           │               │               │              │                 │               │
     │              │──[require]──►│               │               │              │                 │               │
     │              │              │──[new]────────►│              │              │                 │               │
     │              │              │               │──[index($url)]►│             │                 │               │
     │              │              │               │               │─verificar()─►│                 │               │
     │              │              │               │               │◄─OK──────────│                 │               │
     │              │              │               │─[new Orcamento]──────────────►│               │               │
     │              │              │               │              │               │──[new AdmsTecnico]──────────────►│
     │              │              │               │              │               │                 │──[abrirOrcamento()]►│
     │              │              │               │              │               │                 │◄──[INSERT OK]──│
     │              │              │               │              │               │◄──retorno───────│               │
     │              │              │               │              │               │──[dadosOrcamento()]──────────────►│
     │              │              │               │              │               │◄──[SELECT result]──────────────│
     │              │              │               │              │               │──[ConfigView::renderizar()]     │
     │◄─────[HTML renderizado]─────────────────────────────────────────────────────────────────────                │
"""
p = doc.add_paragraph()
p.add_run(seq_text).font.name = "Courier New"
p.runs[0].font.size = Pt(7.5)

# ── 3.8 Diagrama de Actividades ───────────────────────────────────────────────
heading("3.8 Diagrama de Actividades — Fluxo de Atendimento", level=2, size=13)
body(
    "O diagrama de actividades descreve o fluxo completo de atendimento de um equipamento "
    "na oficina, desde a entrada até à entrega ao cliente."
)
act_text = """
 ●  INÍCIO
 │
 ▼
 [Recepcionista regista entrada do equipamento]
 │  → Tabela: entrada_equipamento
 │
 ▼
 [Recepcionista regista cliente e equipamento]
 │  → Tabelas: clientes, equipamento (estado: Recebido)
 │
 ▼
 [Técnico elabora orçamento]
 │  → Tabela: orcamentos (status: Aberto, tipo: Orçamento)
 │  → Adiciona peças necessárias → orc_prod
 │
 ▼
 [Recepcionista apresenta orçamento ao cliente]
 │
 ◇  Cliente aprova?
 │
 │── NÃO ──► [Orçamento rejeitado — equipamento devolvido]  ──► ■ FIM
 │
 └── SIM ──►
 │
 ▼
 [Estado: orcamentos.status = "Aprovado"]
 │  → Tabela: conntas_areceber (registo de adiantamento)
 │
 ▼
 [Técnico executa a reparação]
 │  → Equipamento: estado = "Em Reparação"
 │
 ▼
 [Técnico conclui e actualiza estado]
 │  → Equipamento: estado = "Concluído"
 │  → orcamentos.status = "Concluído"
 │  → comissao: registo automático (VALOR_COMISSAO × valor serviço)
 │  → movimentacao: entrada de caixa
 │  → conntas_areceber.pago = "sim"
 │
 ▼
 [Recepcionista efectua entrega ao cliente]
 │  → Equipamento: estado = "Entregue"
 │
 ▼
 ■  FIM
"""
p = doc.add_paragraph()
p.add_run(act_text).font.name = "Courier New"
p.runs[0].font.size = Pt(9)

# ── 3.9 Funcionalidades por Perfil ────────────────────────────────────────────
heading("3.9 Funcionalidades por Perfil de Utilizador", level=2, size=13)
body(
    "O sistema implementa controlo de acesso baseado em perfil (RBAC — Role-Based Access Control). "
    "O nível de cada utilizador é armazenado no campo nivel da tabela usuario e verificado "
    "em cada sessão autenticada. As secções seguintes detalham as funcionalidades específicas "
    "de cada perfil."
)

heading("3.9.1 Perfil Gerente (adimin)", level=2, size=12)
body(
    "O Gerente possui acesso irrestrito a todos os módulos do sistema. O seu dashboard "
    "(painel inicial) apresenta doze cartões de indicadores em tempo real, organizados em "
    "três linhas:"
)
bullet("Linha financeira: Entrada do Dia, Saída do Dia, Saldo do Dia e Saldo do Mês — "
       "calculados a partir da tabela movimentacao.")
bullet("Linha operacional: Orçamentos Concluídos, Orçamentos Pendentes, Orçamentos Aprovados "
       "e Serviços Pendentes — calculados a partir da tabela orcamentos.")
bullet("Linha de cadastro: Produtos Cadastrados, Total de Clientes, Total de Técnicos "
       "e Total de Recepcionistas — contagens directas das tabelas correspondentes.")

body("Os módulos exclusivos ou de acesso privilegiado do Gerente incluem:")
add_table(
    ["Módulo", "Localização (URL)", "Funcionalidade"],
    [
        ["Gestão de Técnicos", "?url=tecnico", "CRUD completo de técnicos; envio de e-mail de boas-vindas com PHPMailer."],
        ["Gestão de Recepcionistas", "?url=recepcionista", "CRUD completo de recepcionistas."],
        ["Gestão de Fornecedores", "?url=fornecedor", "CRUD de fornecedores com NBI/NIF e tipo de pessoa (Singular/Colectiva)."],
        ["Gestão de Produtos", "?url=produto", "CRUD de peças/componentes com preço de compra, venda e stock actual."],
        ["Categorias", "?url=categoria", "Criação e gestão de categorias de produtos."],
        ["Tipo de Serviço", "?url=tipoServico", "Definição de tipos de serviço com valor padrão."],
        ["Estoque Baixo", "?url=estoque", "Listagem de produtos abaixo do nível mínimo (NIVEL_STOQUE do .env)."],
        ["Compras", "?url=compras", "Registo de compras a fornecedores; actualiza stock e cria conta a pagar."],
        ["Vendas", "?url=vendas", "Registo de vendas de peças; actualiza stock e movimentação."],
        ["Movimentação", "?url=movimentacao", "Visualização e registo do fluxo de caixa (Entrada / Saída)."],
        ["Consultas", "?url=consultas", "Consultas transversais: serviços por técnico, equipamentos por cliente."],
        ["Relatórios", "?url=dashboard", "Relatórios filtráveis por período: serviços, orçamentos, movimentação, compras, vendas, contas, comissões — com exportação PDF via mPDF."],
        ["Gráficos", "?url=graficos", "Gráficos de desempenho financeiro e operacional."],
        ["Chat", "?url=chat", "Chat interno entre utilizadores do sistema."],
    ],
    col_widths=[1.5, 1.7, 3.1]
)

heading("3.9.2 Perfil Técnico (tecnico)", level=2, size=12)
body(
    "O Técnico tem acesso a um painel focado nas suas actividades de reparação e diagnóstico. "
    "O dashboard do Técnico apresenta quatro indicadores: Serviços Concluídos, Orçamentos "
    "Abertos (a ele atribuídos), Comissões do Dia e Comissões do Mês. "
    "Abaixo dos indicadores é apresentada uma lista de Serviços Pendentes."
)
add_table(
    ["Módulo", "URL", "Funcionalidade"],
    [
        ["Orçamentos", "?url=orcamento", "Criar, editar e eliminar orçamentos de serviço; adicionar peças (orc_prod); imprimir PDF do orçamento."],
        ["Serviços", "?url=servico", "Visualizar os serviços atribuídos ao técnico autenticado; actualizar estado do equipamento; registar diagnóstico técnico."],
        ["Comissões", "?url=comissoes", "Consultar o histórico de comissões recebidas, filtradas por período."],
        ["Relatório Comissão", "?url=relatorioTecnico", "Gerar relatório PDF das comissões do técnico."],
        ["Perfil", "?url=perfil", "Editar dados pessoais e foto de perfil."],
        ["Chat", "?url=chat", "Chat interno com outros utilizadores."],
    ],
    col_widths=[1.5, 1.5, 3.3]
)
body(
    "O cálculo da comissão é automático: ao marcar um orçamento como Concluído, o sistema "
    "multiplica o valor do serviço pela constante VALOR_COMISSAO (por omissão 0.30, i.e. 30%) "
    "e insere um registo na tabela comissao com o NIF do técnico. Esta percentagem é "
    "configurável no ficheiro .env sem necessidade de alterar código."
)

heading("3.9.3 Perfil Recepcionista (recep)", level=2, size=12)
body(
    "A Recepcionista gere o front-office da oficina: recebe os equipamentos, interage com os "
    "clientes e controla o fluxo financeiro do dia-a-dia. O seu dashboard apresenta os "
    "mesmos oito indicadores financeiros e operacionais do Gerente, mas sem acesso aos "
    "módulos de cadastro de pessoal e relatórios completos."
)
add_table(
    ["Módulo", "URL", "Funcionalidade"],
    [
        ["Clientes", "?url=cliente", "CRUD completo de clientes: NBI, NIF, nome, contacto, morada."],
        ["Equipamentos", "?url=equipamento", "Registo de equipamentos associados a clientes; actualização de estado; registo de defeito reportado e diagnóstico."],
        ["Entrada de Equipamentos", "?url=entradaEquipamento", "Registo formal de entrada de equipamentos na oficina, com data e técnico responsável."],
        ["Orçamentos (Recepção)", "?url=orcamentoRecepcao", "Visualizar, aprovar e gerir orçamentos submetidos pelos técnicos."],
        ["Contas a Pagar", "?url=contasPagar", "Gerir contas a pagar (pagamento a fornecedores, despesas); marcar como pago."],
        ["Conta a Receber", "?url=contaReceber", "Gerir contas a receber (pagamentos de clientes); marcar como pago."],
        ["Movimentação", "?url=movimentacao", "Consultar movimentação de caixa do dia/período."],
        ["Compras", "?url=compras", "Registar compras a fornecedores (com permissão concedida pelo gerente)."],
        ["Consultas", "?url=consultas", "Consultas rápidas: histórico de equipamentos por cliente, estado de serviços."],
        ["Relatórios", "?url=relatorio", "Relatórios disponíveis para o perfil recepcionista."],
        ["Perfil / Foto", "?url=perfil / ?url=editarFoto", "Gestão do perfil pessoal."],
        ["Chat", "?url=chat", "Comunicação interna com técnicos e gerente."],
    ],
    col_widths=[1.8, 1.7, 2.8]
)

page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  CAP 4 — CONCLUSÕES
# ═══════════════════════════════════════════════════════════════════════════════
heading("4. CONCLUSÕES E TRABALHO FUTURO", level=1, size=14)
body(
    "O sistema de gestão de assistência técnica informática desenvolvido cumpre integralmente "
    "os objectivos traçados no início do projecto. A aplicação MVC em PHP/MySQL provou ser "
    "uma solução robusta, económica (sem custos de licenciamento) e de fácil manutenção, "
    "adequada ao contexto de uma PME angolana do sector de informática."
)
body(
    "Os principais ganhos obtidos com a implementação do sistema são: (i) eliminação dos "
    "registos em papel e a consequente redução de erros de transcrição; (ii) rastreabilidade "
    "completa de cada equipamento desde a entrada até à entrega; (iii) cálculo automático "
    "de comissões, eliminando disputas entre a gestão e os técnicos; (iv) visibilidade "
    "financeira em tempo real através do dashboard; (v) geração automática de relatórios "
    "e documentos PDF para arquivo e entrega ao cliente."
)
body(
    "Como trabalho futuro, identificam-se as seguintes melhorias de maior impacto:"
)
bullet("Migração do hash de senha de MD5 para bcrypt ou Argon2, eliminando a vulnerabilidade "
       "mais crítica do sistema actual.")
bullet("Desenvolvimento de uma API REST para permitir a criação de uma aplicação móvel "
       "(Android/iOS) que permita aos técnicos actualizarem o estado dos serviços em campo.")
bullet("Implementação de notificações automáticas por SMS ao cliente quando o seu equipamento "
       "fica pronto para levantamento.")
bullet("Adição de um módulo de orçamento por fotografia, permitindo ao cliente enviar fotos "
       "do equipamento para pré-diagnóstico remoto.")
bullet("Substituição das queries SQL directas por um ORM (e.g. Eloquent / Doctrine) para "
       "melhorar a manutenibilidade e eliminar riscos de SQL injection residuais.")
bullet("Implementação de testes automáticos (PHPUnit) para garantir a regressão em futuras "
       "iterações de desenvolvimento.")

page_break()

# ═══════════════════════════════════════════════════════════════════════════════
#  REFERÊNCIAS
# ═══════════════════════════════════════════════════════════════════════════════
heading("REFERÊNCIAS BIBLIOGRÁFICAS", level=1, size=14)

refs = [
    "GAMMA, E.; HELM, R.; JOHNSON, R.; VLISSIDES, J. Design Patterns: Elements of Reusable Object-Oriented Software. Addison-Wesley, 1994.",
    "LAUDON, K. C.; LAUDON, J. P. Management Information Systems: Managing the Digital Firm. 16.ª ed. Pearson Education, 2021.",
    "PRESSMAN, R. S.; MAXIM, B. R. Engenharia de Software: Uma Abordagem Profissional. 8.ª ed. McGraw-Hill, 2016.",
    "SOMMERVILLE, I. Software Engineering. 10.ª ed. Pearson Education, 2019.",
    "TURBAN, E.; VOLONINO, L.; WOOD, G. Information Technology for Management: Driving Digital Transformation to Increase Local and Global Performance, Growth and Sustainability. 11.ª ed. Wiley, 2018.",
    "PHP GROUP. PHP Manual. Disponível em: https://www.php.net/manual/. Acesso em: junho de 2026.",
    "MYSQL. MySQL 8.0 Reference Manual. Oracle Corporation. Disponível em: https://dev.mysql.com/doc/refman/8.0/en/. Acesso em: junho de 2026.",
    "OMG — OBJECT MANAGEMENT GROUP. Unified Modeling Language Specification, version 2.5.1. 2017. Disponível em: https://www.omg.org/spec/UML/.",
    "BOOTSTRAP. Bootstrap 4 Documentation. Disponível em: https://getbootstrap.com/docs/4.6/. Acesso em: junho de 2026.",
    "SYNACTIS. mPDF Documentation. Disponível em: https://mpdf.github.io/. Acesso em: junho de 2026.",
    "PHPMAILER. PHPMailer Documentation. Disponível em: https://github.com/PHPMailer/PHPMailer. Acesso em: junho de 2026.",
    "COMPOSER. Dependency Manager for PHP. Disponível em: https://getcomposer.org/. Acesso em: junho de 2026.",
]
for i, ref in enumerate(refs, 1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(f"{i}. {ref}")
    r.font.size = Pt(11)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(1.25)
    p.paragraph_format.first_line_indent = Cm(-1.25)

# ═══════════════════════════════════════════════════════════════════════════════
#  GRAVAR
# ═══════════════════════════════════════════════════════════════════════════════
output_path = r"c:\xampp\htdocs\oficina-de-informatica\Monografia_Assistencia_Tecnica_Informatica.docx"
doc.save(output_path)
print(f"Monografia gerada com sucesso: {output_path}")
